"""无需模型或 API key；真实 Chroma 测试需要安装 chromadb。

python -m unittest discover -s tests -v
"""
import importlib.util
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import unittest
from pathlib import Path
from unittest.mock import patch

import ingestion.clean_chunk as clean_chunk
import ingestion.embed_store as embed_store
import ingestion.source_store as source_store
from utils.config_loader import settings
from utils.index_audit import assert_index_collection, validate_storage_paths
from utils.schema import CanonicalBlock
from ingestion.source_store import SourceResolver, SourceResolutionError, located_blocks


class SourceTests(unittest.TestCase):
    def setUp(self):
        child_root = os.getenv("RAG_TEST_SUBPROCESS_ROOT")
        if child_root:
            self.root = Path(child_root)
        else:
            self.temp = tempfile.TemporaryDirectory()
            self.addCleanup(self.temp.cleanup)
            self.root = Path(self.temp.name)
        self.kb = self.root / "kb"
        (self.kb / "topology").mkdir(parents=True, exist_ok=True)
        self.store = self.root / "source"
        self.patch = patch.object(settings.paths, "source_store_dir", self.store)
        self.patch.start()
        self.addCleanup(self.patch.stop)

    def build(self, name="guide.txt", text="BUCK 的电感纹波需要根据开关频率计算。\n完整原文不能出现在向量数据库里。", size=40):
        path = self.kb / "topology" / name
        path.write_text(text, encoding="utf-8")
        return clean_chunk.build_chunks_for_file(path, self.kb, size)

    def test_jsonl_has_only_index_and_roundtrips_full_content(self):
        chunks = self.build()
        output = self.root / "index.jsonl"
        output.write_text("\n".join(json.dumps(c.to_index_dict(), ensure_ascii=False) for c in chunks), encoding="utf-8")
        loaded = embed_store.load_chunks(output)
        self.assertEqual([c.content for c in loaded], [c.content for c in chunks])
        for chunk in chunks:
            index = chunk.to_index_dict()
            self.assertFalse({"content", "document", "documents", "title", "section_path", "tags"} & set(index))
            self.assertTrue(json.loads(index["keywords"]))
            item = SourceResolver().hydrate({"chunk_id": chunk.chunk_id, "metadata": chunk.chroma_metadata()})
            self.assertEqual(item["document"], chunk.content)
            self.assertEqual(item["metadata"]["section_path"], chunk.section_path)

    def test_lines_columns_pages_and_repeated_text_are_exact(self):
        blocks = [CanonicalBlock("text", "第一行\n" + "A" * 27 + "\n重复行\n重复行", section_path="同一章节", page=1),
                  CanonicalBlock("text", "第二页\n最后一行", section_path="同一章节", page=2)]
        units, slices = located_blocks(blocks, "title", 11)
        self.assertEqual(len(units), 2)
        for unit_id, text in units.items():
            self.assertEqual("".join(block.content for block, loc in slices if loc["source_unit"] == unit_id), text)
        for block, loc in slices:
            text = units[loc["source_unit"]]
            self.assertEqual(block.content, text[loc["char_start"]:loc["char_end"]])
            self.assertLessEqual(len(block.content), 11)
            self.assertEqual(loc["line_start"], text[:loc["char_start"]].count("\n") + 1)
            self.assertEqual(loc["line_end"], text[:loc["char_end"] - 1].count("\n") + 1)
        self.assertTrue(any(loc["column_start"] > 1 for _, loc in slices))
        self.assertEqual({block.page for block, _ in slices}, {1, 2})

    def test_all_format_location_fields_survive(self):
        blocks = [CanonicalBlock("table_row", "单元格1\n单元格2", sheet_name="Sheet A", row_start=3, row_end=3),
                  CanonicalBlock("slide_text", "幻灯片完整内容", slide_number=2),
                  CanonicalBlock("image_ocr", "图片文字完整内容", image_index=1),
                  CanonicalBlock("transcript", "音频转录完整内容", media_start_sec=2.0, media_end_sec=5.0)]
        _, slices = located_blocks(blocks, "title", 100)
        for original, (result, _) in zip(blocks, slices):
            for field in ("sheet_name", "row_start", "row_end", "slide_number", "image_index", "media_start_sec", "media_end_sec"):
                self.assertEqual(getattr(original, field), getattr(result, field))

    def test_tampered_missing_or_wrong_locator_fails_closed(self):
        chunk = self.build()[0]
        item = {"chunk_id": chunk.chunk_id, "metadata": chunk.chroma_metadata()}
        for key, value in (("char_start", 999), ("domain", "device"), ("line_start", 999), ("source_ref", "../../secret")):
            with self.subTest(key=key), self.assertRaises(SourceResolutionError):
                SourceResolver().hydrate({**item, "metadata": {**item["metadata"], key: value}})
        snapshot = self.store / chunk.source_ref[:2] / (chunk.source_ref + ".json")
        snapshot.write_text("tampered", encoding="utf-8")
        with self.assertRaises(SourceResolutionError):
            SourceResolver().hydrate(item)
        snapshot.unlink()
        with self.assertRaises(SourceResolutionError):
            SourceResolver().hydrate(item)

    def test_source_versions_do_not_silently_change_old_results(self):
        old = self.build()[0]
        new = self.build(text="更新之后的原文已经完全不同。\n只有重建索引才应检索新版本。")[0]
        self.assertNotEqual(old.source_ref, new.source_ref)
        for chunk in (old, new):
            self.assertEqual(SourceResolver().hydrate({"chunk_id": chunk.chunk_id, "metadata": chunk.chroma_metadata()})["document"], chunk.content)

    def test_short_content_is_not_lost(self):
        # Docling 路径会前缀确定性上下文头（CCH）；短正文必须完整保留在正文部分。
        content = self.build(text="短句")[0].content
        self.assertTrue(content.endswith("短句"))

    def test_old_plaintext_jsonl_is_rejected(self):
        path = self.root / "old.jsonl"
        path.write_text(json.dumps(self.build()[0].to_dict()), encoding="utf-8")
        with self.assertRaisesRegex(RuntimeError, "旧版正文"):
            embed_store.load_chunks(path)

    def test_parser_failure_preserves_previous_manifest_and_does_not_log_body(self):
        self.build()
        output, error = self.root / "index.jsonl", self.root / "errors.jsonl"
        output.write_text("previous index", encoding="utf-8")
        with patch.object(clean_chunk, "build_chunks_for_file", side_effect=ValueError("SECRET BODY")):
            with self.assertRaises(RuntimeError):
                clean_chunk.run(self.kb, output, error, 100)
        self.assertEqual(output.read_text(encoding="utf-8"), "previous index")
        self.assertNotIn("SECRET BODY", error.read_text(encoding="utf-8"))

    def test_source_store_must_be_outside_chroma(self):
        for chroma, source in ((self.root, self.root / "source"), (self.root / "db", self.root), (self.root, self.root)):
            with self.assertRaises(ValueError):
                validate_storage_paths(chroma, source)
        validate_storage_paths(self.root / "db", self.store)

    @unittest.skipUnless(importlib.util.find_spec("chromadb"), "需要 chromadb")
    def test_real_chroma_retrieval_expansion_reingest_and_audit(self):
        # Chroma 的 Rust/SQLite 句柄在 Windows 上可能持有到进程退出。
        # 子进程跑真实持久库断言，退出后父进程再删除临时数据。
        if not os.getenv("RAG_TEST_SUBPROCESS_ROOT"):
            result = subprocess.run(
                [sys.executable, "-X", "utf8", "-m", "unittest", "discover",
                 "-s", "tests", "-p", "test_index_only.py", "-k", "real_chroma_retrieval"],
                cwd=Path(__file__).resolve().parents[1],
                env={**os.environ, "RAG_TEST_SUBPROCESS_ROOT": str(self.root)},
                capture_output=True, encoding="utf-8", timeout=90,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            return
        import chromadb
        import numpy as np
        from pipeline.retrieval import retriever as search
        from pipeline.context import expander
        from pipeline.generation.generator import _build_grounded_context

        text = "# 原理\nBUCK parent unique sentence.\n## 计算\nBUCK 电感需要根据压差和频率计算。\n" + "完整长行原文不应存入向量库。" * 9
        chunks = self.build("guide.md", text, size=45)
        client = chromadb.PersistentClient(path=str(self.root / "chroma"))
        collection = embed_store.get_collection(client, "topology")
        vectors = np.array([[1.0, 0.1, 0.0]] * len(chunks))
        embed_store.replace_domain_documents(collection, chunks, vectors)
        self.assertEqual(assert_index_collection(collection, scan=True), len(chunks))
        raw = collection.get(include=["documents", "metadatas", "embeddings"])
        self.assertTrue(all(d is None for d in raw["documents"]))
        self.assertEqual(len(raw["embeddings"]), len(chunks))
        self.assertNotIn("完整长行原文不应存入向量库。", json.dumps(raw["metadatas"], ensure_ascii=False))

        found = search.query_collection(collection=collection, query_embedding=vectors[0], candidate_limit=20)
        self.assertEqual({i["document"] for i in found}, {c.content for c in chunks})
        anchor = next(i for i in found if " / " in i["metadata"]["section_path"])
        section = expander.expand_same_section(client, anchor)
        self.assertTrue(section)
        self.assertTrue(all(i["document"] for i in section))
        parent = expander.expand_parent_section(client, anchor)
        self.assertTrue(parent)
        self.assertIn("parent unique", parent[0]["document"])
        neighbors = expander.expand_neighbors(client, anchor, before=1, after=1)
        self.assertTrue(neighbors)
        class Model:
            def encode(self, texts, **kwargs):
                return np.array([[1.0, 0.1, 0.0]] * len(texts))
        targeted = expander.expand_targeted_search(client=client, embedding_model=Model(), query="BUCK", domains=["topology"])
        self.assertTrue(targeted)
        self.assertTrue(all(i["document"] for i in targeted))
        prompt, sources = _build_grounded_context(found)
        self.assertIn(found[0]["document"], prompt)
        self.assertEqual(len(sources), len(found))

        # 完整入口联调：真实解析/快照/索引/Chroma；替换模型避免下载和外部调用。
        from ingestion import ingest
        manifest = self.root / "processed" / "index.jsonl"
        with patch.object(settings.paths, "kb_root", self.kb), \
             patch.object(settings.paths, "chunk_jsonl", manifest), \
             patch.object(settings.paths, "error_jsonl", self.root / "errors.jsonl"), \
             patch.object(settings.ingestion.chunking, "section_max_chars", 45), \
             patch.object(settings.paths, "chroma_dir", self.root / "chroma"), \
             patch.object(embed_store, "get_client", return_value=client), \
             patch.object(embed_store, "load_embedding_model", return_value=Model()):
            ingest.main([])
        self.assertNotIn('"content":', manifest.read_text(encoding="utf-8"))
        self.assertEqual(collection.count(), len(chunks))
        def choose(**kwargs):
            return {"sufficient": True, "reason": "test", "expansion_requests": []}, kwargs["candidates"]
        def answer(**kwargs):
            content, references = _build_grounded_context(kwargs["selected_contexts"])
            return content, references
        with patch.object(search, "get_client", return_value=client), \
             patch.object(search, "load_embedding_model", return_value=Model()), \
             patch.object(search, "create_llm_client", return_value=object()), \
             patch.object(settings.rag.generation, "enabled", True), \
             patch.object(settings.rag.retrieval.rerank, "enabled", False), \
             patch.object(settings.rag.retrieval.hybrid, "enabled", False), \
             patch.object(settings.rag.retrieval.multi_hop, "enabled", False), \
             patch.object(search, "analyze_query", return_value={"domains": ["topology"], "keywords": ["BUCK"]}), \
             patch.object(search, "select_contexts", side_effect=choose), \
             patch.object(search, "generate_answer", side_effect=answer):
            response = search.run("BUCK calculation")
        self.assertTrue(response["selected_contexts"])
        self.assertIn("lines=", response["answer_sources"][0]["location"])

        # 替换同文档后旧分块必须消失，当前记录依然无正文。
        updated = self.build("guide.md", "# 原理\n更新的正文只在本地原文目录中。", size=45)
        embed_store.replace_domain_documents(collection, updated, vectors[:len(updated)])
        self.assertEqual(collection.count(), len(updated))
        self.assertEqual(assert_index_collection(collection, scan=True), len(updated))
        embed_store.replace_domain_documents(collection, updated, vectors[:len(updated)])
        self.assertEqual(collection.count(), len(updated))

        # 有版本标记也不能绕过正文审计。
        collection.upsert(ids=["bad"], documents=["SECRET BODY"], embeddings=[[1., 0., 0.]], metadatas=[updated[0].chroma_metadata()])
        with self.assertRaisesRegex(RuntimeError, "documents"):
            assert_index_collection(collection, scan=True)
        with self.assertRaises(RuntimeError):
            embed_store.replace_domain_documents(collection, updated, vectors[:len(updated)])
        legacy = client.create_collection("legacy", embedding_function=None)
        with self.assertRaisesRegex(RuntimeError, "旧版"):
            assert_index_collection(legacy)

    @unittest.skipUnless(all(importlib.util.find_spec(m) for m in ("docx", "pptx", "openpyxl", "bs4", "lxml")), "需要 Office/HTML 解析器依赖")
    def test_real_office_html_csv_files_roundtrip(self):
        from docx import Document
        from pptx import Presentation
        from openpyxl import Workbook

        directory = self.kb / "topology"
        doc = Document()
        doc.add_heading("Design", level=1)
        doc.add_paragraph("BUCK original paragraph with formula L = 10 uH.")
        doc.save(directory / "guide.docx")
        slides = Presentation()
        slide = slides.slides.add_slide(slides.slide_layouts[1])
        slide.shapes.title.text = "BUCK Design"
        slide.placeholders[1].text = "Original slide content and frequency details."
        slides.save(directory / "guide.pptx")
        book = Workbook()
        book.active.title = "Components"
        book.active.append(["Name", "Value"])
        book.active.append(["Inductor", "10 uH"])
        book.save(directory / "guide.xlsx")
        book.close()
        (directory / "guide.csv").write_text("Name,Value\nInductor,10 uH\nCapacitor,20 uF\n", encoding="utf-8")
        (directory / "guide.html").write_text("<h1>BUCK</h1><p>Original HTML paragraph with frequency details.</p>", encoding="utf-8")
        for path in directory.iterdir():
            with self.subTest(format=path.suffix):
                chunks = clean_chunk.build_chunks_for_file(path, self.kb, 30)
                self.assertTrue(chunks)
                for chunk in chunks:
                    restored = SourceResolver().hydrate({"chunk_id": chunk.chunk_id, "metadata": chunk.chroma_metadata()})
                    self.assertEqual(restored["document"], chunk.content)
                    self.assertIsNone(chunk.page)
                    if path.suffix in (".xlsx", ".csv"):
                        self.assertGreaterEqual(chunk.row_start, 2)
                        self.assertEqual(chunk.row_start, chunk.row_end)
                    if path.suffix == ".pptx":
                        self.assertEqual(chunk.slide_number, 1)

    @unittest.skipUnless(importlib.util.find_spec("fitz"), "需要 PyMuPDF")
    def test_real_pdf_pages_remain_separate(self):
        import fitz
        path = self.kb / "topology" / "two_pages.pdf"
        doc = fitz.open()
        for number in (1, 2):
            page = doc.new_page()
            page.insert_text((72, 72), f"Page {number} first source line.\nSecond source line with BUCK calculation and frequency details.")
        doc.save(path)
        doc.close()
        chunks = clean_chunk.build_chunks_for_file(path, self.kb, 1000)
        self.assertEqual({c.page for c in chunks}, {1, 2})
        for chunk in chunks:
            self.assertEqual(chunk.line_start, 1)
            self.assertEqual(chunk.line_end, 2)
            self.assertEqual(SourceResolver().hydrate({"chunk_id": chunk.chunk_id, "metadata": chunk.chroma_metadata()})["document"], chunk.content)


class IncrementalIngestTests(unittest.TestCase):
    """增量切片：只重解析变化的文件，并让索引始终等于「当前知识库的完整切片」。"""

    SIZE = 45

    def setUp(self):
        child_root = os.getenv("RAG_TEST_SUBPROCESS_ROOT")
        if child_root:
            self.root = Path(child_root)
        else:
            self.temp = tempfile.TemporaryDirectory()
            self.addCleanup(self.temp.cleanup)
            self.root = Path(self.temp.name)

        self.kb = self.root / "kb"
        (self.kb / "topology").mkdir(parents=True, exist_ok=True)
        self.store = self.root / "source"
        self.output = self.root / "index.jsonl"
        self.errors = self.root / "errors.jsonl"

        for target, name, value in (
            (settings.paths, "source_store_dir", self.store),
            (settings.paths, "kb_root", self.kb),
            (settings.paths, "chunk_jsonl", self.output),
            (settings.paths, "error_jsonl", self.errors),
            (settings.ingestion.chunking, "section_max_chars", self.SIZE),
            # 测试用的知识库只有两三个文件，默认阈值下删一个就超限。
            # 阈值行为由 test_incremental_refuses_mass_delete_without_force 单独覆盖。
            (settings.ingestion.incremental, "max_delete_ratio", 1.0),
        ):
            patcher = patch.object(target, name, value)
            patcher.start()
            self.addCleanup(patcher.stop)

    # ---------- 工具 ----------

    def write(self, name, text):
        path = self.kb / "topology" / name
        path.write_text(text, encoding="utf-8")
        return path

    def remove(self, name):
        (self.kb / "topology" / name).unlink()

    def full(self):
        return clean_chunk.run(self.kb, self.output, self.errors, self.SIZE)

    def incremental(self, **kwargs):
        return clean_chunk.run(
            self.kb, self.output, self.errors, self.SIZE, incremental=True, **kwargs
        )

    def resume(self, **kwargs):
        return clean_chunk.run(
            self.kb, self.output, self.errors, self.SIZE, resume=True, **kwargs
        )

    def tmp_index(self):
        return self.output.with_suffix(self.output.suffix + ".tmp")

    def tmp_state_file(self):
        from ingestion.index_state import state_path

        return state_path(self.tmp_index())

    def tmp_completeness(self):
        return json.loads(self.tmp_state_file().read_text(encoding="utf-8"))["complete"]

    def _patch_parse_failure(self, failing_name):
        """让指定文件解析失败，其余正常——用来模拟「中途某个文件坏了」的崩溃现场。"""
        original = clean_chunk.build_chunks_for_file

        def side_effect(*, path, kb_root, section_max_chars):
            if Path(path).name == failing_name:
                raise ValueError("模拟解析失败")
            return original(path=path, kb_root=kb_root, section_max_chars=section_max_chars)

        return patch.object(clean_chunk, "build_chunks_for_file", side_effect=side_effect)

    def rows_by_path(self, path=None):
        rows = {}
        for line in (path or self.output).read_text(encoding="utf-8").splitlines():
            if line.strip():
                rows.setdefault(json.loads(line)["relative_path"], []).append(line)
        return rows

    def state_file(self):
        from ingestion.index_state import state_path

        return state_path(self.output)

    def tree_stamp(self, *paths):
        stamp = {}
        for base in paths:
            base = Path(base)
            targets = [base] if base.is_file() else sorted(p for p in base.rglob("*") if p.is_file())
            if not base.exists():
                stamp[str(base)] = None
                continue
            for path in targets:
                stat = path.stat()
                stamp[str(path)] = (stat.st_mtime_ns, stat.st_size)
        return stamp

    def assert_equals_full_rebuild(self):
        """增量产出的索引必须与「在干净目录上跑全量」逐字节一致。"""
        reference_kb = self.root / "reference_kb"
        if reference_kb.exists():
            shutil.rmtree(reference_kb)
        shutil.copytree(self.kb, reference_kb)
        reference_index = self.root / "reference_index.jsonl"
        with patch.object(settings.paths, "kb_root", reference_kb):
            clean_chunk.run(
                reference_kb, reference_index, self.root / "reference_errors.jsonl", self.SIZE
            )
        self.assertEqual(self.output.read_bytes(), reference_index.read_bytes())

    # ---------- 核心行为 ----------

    def test_incremental_no_change_is_byte_identical_and_parses_nothing(self):
        self.write("a.md", "# 原理\nBUCK 电感需要根据压差和频率计算。")
        self.write("b.md", "# 计算\n电容纹波与开关频率有关。")
        self.full()

        before = self.output.read_bytes()
        stamp = self.output.stat().st_mtime_ns

        with patch.object(
            clean_chunk, "build_chunks_for_file", side_effect=AssertionError("不应重新解析")
        ):
            summary = self.incremental()

        self.assertEqual(summary["unchanged"], 2)
        self.assertEqual(summary["chunks"], 0)
        # 索引文件连 mtime 都不能变，否则长驻服务的词法索引会白重建一次。
        self.assertEqual(self.output.stat().st_mtime_ns, stamp)
        self.assertEqual(self.output.read_bytes(), before)

    def test_incremental_reparses_only_changed_file(self):
        self.write("a.md", "# 原理\nBUCK 电感需要根据压差和频率计算。")
        self.write("b.md", "# 计算\n电容纹波与开关频率有关。")
        self.full()
        before = self.rows_by_path()

        self.write("b.md", "# 计算\n完全换过的正文，只有电容纹波这一句留下。")

        original = clean_chunk.build_chunks_for_file
        parsed = []

        def spy(**kwargs):
            parsed.append(Path(kwargs["path"]).name)
            return original(**kwargs)

        with patch.object(clean_chunk, "build_chunks_for_file", side_effect=spy):
            summary = self.incremental()

        self.assertEqual(parsed, ["b.md"])
        self.assertEqual(summary["changed"], 1)
        self.assertEqual(summary["unchanged"], 1)
        after = self.rows_by_path()
        self.assertEqual(after["topology/a.md"], before["topology/a.md"])
        self.assertNotEqual(after["topology/b.md"], before["topology/b.md"])
        self.assert_equals_full_rebuild()

    def test_incremental_removes_rows_of_deleted_file(self):
        self.write("a.md", "# 原理\nBUCK 电感需要根据压差和频率计算。")
        self.write("b.md", "# 计算\n电容纹波与开关频率有关。")
        self.full()
        before = self.rows_by_path()

        self.remove("b.md")
        summary = self.incremental()

        self.assertEqual(summary["deleted"], 1)
        after = self.rows_by_path()
        self.assertNotIn("topology/b.md", after)
        self.assertEqual(after["topology/a.md"], before["topology/a.md"])
        self.assert_equals_full_rebuild()

    def test_incremental_rename_is_delete_plus_add(self):
        self.write("a.md", "# 原理\nBUCK 电感需要根据压差和频率计算。")
        self.write("b.md", "# 计算\n电容纹波与开关频率有关。")
        self.full()
        old_document_ids = {
            json.loads(line)["document_id"] for line in self.output.read_text(encoding="utf-8").splitlines() if line.strip()
        }

        self.write("renamed.md", (self.kb / "topology" / "b.md").read_text(encoding="utf-8"))
        self.remove("b.md")
        self.incremental()

        after = self.rows_by_path()
        self.assertNotIn("topology/b.md", after)
        self.assertIn("topology/renamed.md", after)
        new_ids = {json.loads(line)["document_id"] for rows in after.values() for line in rows}
        self.assertNotEqual(new_ids, old_document_ids)
        self.assert_equals_full_rebuild()

    def test_incremental_equals_full_rebuild_across_many_rounds(self):
        self.write("a.md", "# 原理\nBUCK 电感需要根据压差和频率计算。\n## 计算\n电感纹波与频率有关。")
        self.write("b.md", "# 布局\nPCB 走线宽度影响温升。")
        self.write("c.md", "# 器件\nMOSFET 导通损耗需要按电流估算。")
        self.full()

        self.write("d.md", "# 控制\n环路补偿要保证相位裕度。")          # 新增
        self.write("b.md", "# 布局\n反复改过的走线宽度与铜厚都会影响温升。")  # 修改
        self.remove("c.md")                                              # 删除
        self.incremental()
        self.assert_equals_full_rebuild()

        self.write("e.md", "# 拓扑\n反激变压器需要按匝比设计。")
        self.remove("d.md")
        self.incremental()
        self.assert_equals_full_rebuild()

    def _write_partial_twins(self):
        """同目录、同 stem（换扩展名）、部分内容相同的两份文件。

        全量按路径序 a.md 先到，a.txt 里那份相同的块被跨文件去重挤掉；但 a.txt 仍有
        独有的块落盘，因此下次增量会把它判成「未变」——这正是「赢家被删/被改后，
        被挤掉的内容永久静默消失」的场景。
        """
        shared = "## 共用章节\nBUCK 电感需要根据压差和频率计算。\n"
        self.write("a.md", "# 文档\n" + shared + "## 仅md\nMOSFET 导通损耗按电流估算。\n")
        self.write("a.txt", "# 文档\n" + shared + "## 仅txt\n环路补偿要保证相位裕度。\n")

    def test_incremental_stem_twin_regains_dropped_chunks(self):
        self._write_partial_twins()
        self.full()

        rows = self.rows_by_path()
        self.assertEqual(len(rows["topology/a.md"]), 2)
        self.assertEqual(len(rows["topology/a.txt"]), 1)  # 共用块被 a.md 挤掉

        self.remove("a.md")
        summary = self.incremental()

        # a.txt 必须被孪生扩张提升为重新解析，否则被挤掉的块永久消失。
        self.assertEqual(summary["changed"], 1)
        self.assertEqual(summary["added"], 0)
        self.assertEqual(len(self.rows_by_path()["topology/a.txt"]), 2)
        self.assert_equals_full_rebuild()

    def test_incremental_stem_twin_reparsed_when_sibling_changes(self):
        self._write_partial_twins()
        self.full()

        self.write("a.md", "# 文档\n## 换过\n电感换成电容，频率换成占空比。\n")

        original = clean_chunk.build_chunks_for_file
        parsed = []

        def spy(**kwargs):
            parsed.append(Path(kwargs["path"]).name)
            return original(**kwargs)

        with patch.object(clean_chunk, "build_chunks_for_file", side_effect=spy):
            self.incremental()

        self.assertIn("a.txt", parsed)
        self.assert_equals_full_rebuild()

    def test_resume_after_partial_failure_equals_full_rebuild(self):
        # 旧实现按 relative_path 跳过「已完成」的文件，而这里的 a.md 因为排在前面、
        # 解析失败过，所以 a.txt 抢走了两者共用块的去重胜负；续跑时 a.md 反而落败，
        # 产出与全量重跑不同。现在两者共用同一套装配逻辑。
        self._write_partial_twins()
        with self._patch_parse_failure("a.md"):
            with self.assertRaises(RuntimeError):
                self.full()

        self.assertTrue(self.tmp_index().exists())
        self.assertFalse(self.output.exists())  # 有错误 -> 不发布

        summary = self.resume()

        # a.md 重切，a.txt 因孪生扩张一并重切，去重胜负与全量一致
        self.assertEqual(summary["changed"], 1)
        self.assertEqual(len(self.rows_by_path()["topology/a.md"]), 2)
        self.assertEqual(len(self.rows_by_path()["topology/a.txt"]), 1)
        self.assert_equals_full_rebuild()

    def test_resume_reparses_file_changed_during_interruption(self):
        self.write("a.md", "# 原理\nBUCK 电感需要根据压差和频率计算。")
        self.write("b.md", "# 计算\n电容纹波与开关频率有关。")
        with self._patch_parse_failure("b.md"):
            with self.assertRaises(RuntimeError):
                self.full()

        def hashes(text):
            return {
                json.loads(line)["content_hash"]
                for line in text.splitlines()
                if line.strip() and "a.md" in line
            }

        before = hashes(self.tmp_index().read_text(encoding="utf-8"))

        # 中断期间把已经写进 .tmp 的文件改掉：续跑必须重切它，不能当成已完成
        self.write("a.md", "# 原理\n换过的正文：电感换成电容，频率换成占空比。")
        self.resume()

        after = hashes(self.output.read_text(encoding="utf-8"))
        self.assertNotEqual(before, after)
        self.assert_equals_full_rebuild()

    def test_resume_discards_partially_written_last_file(self):
        self.write("a.md", "# 原理\nBUCK 电感需要根据压差和频率计算。")
        self.write("b.md", "# 计算\n电容纹波与开关频率有关。")
        self.full()

        # 模拟原生崩溃：把索引搬成 .tmp，并让最后一个文件的记录只剩半行
        lines = self.output.read_text(encoding="utf-8").splitlines()
        self.tmp_index().write_text(
            "\n".join(lines[:-1] + [lines[-1][:60]]) + "\n", encoding="utf-8"
        )
        self.output.unlink()

        summary = self.resume()

        # 没有 complete 标记 -> 视为原生崩溃，最后一组整组丢弃重做（宁可多切也不错用半截数据）
        self.assertEqual(summary["unchanged"], 0)
        self.assert_equals_full_rebuild()

    def test_full_run_resets_tmp_completeness_marker(self):
        # 上一次异常退出留下了 .tmp 与 complete=true。
        self.write("a.md", "# 原理\nBUCK 电感需要根据压差和频率计算。")
        self.write("b.md", "# 计算\n电容纹波与开关频率有关。")
        with self._patch_parse_failure("b.md"):
            with self.assertRaises(RuntimeError):
                self.full()
        self.assertTrue(self.tmp_completeness())

        # 用户直接重跑全量，进程在中途被原生杀掉。KeyboardInterrupt 不是 Exception，
        # 不会被逐文件捕获，于是循环后的 complete=true 永远不会写。
        original = clean_chunk.build_chunks_for_file
        parsed = []

        def side_effect(*, path, kb_root, section_max_chars):
            parsed.append(Path(path).name)
            if len(parsed) == 2:
                raise KeyboardInterrupt("模拟进程被原生杀掉")
            return original(path=path, kb_root=kb_root, section_max_chars=section_max_chars)

        with patch.object(clean_chunk, "build_chunks_for_file", side_effect=side_effect):
            with self.assertRaises(KeyboardInterrupt):
                self.full()

        # 关键：不能残留上一次的 true，否则这份「被写了一半」的 .tmp 会被当成完整产物
        # 复用，半截数据直接进索引。
        self.assertFalse(self.tmp_completeness())

    def test_resume_removes_file_deleted_during_interruption(self):
        self.write("a.md", "# 原理\nBUCK 电感需要根据压差和频率计算。")
        self.write("b.md", "# 计算\n电容纹波与开关频率有关。")
        with self._patch_parse_failure("b.md"):
            with self.assertRaises(RuntimeError):
                self.full()

        self.remove("a.md")
        summary = self.resume()

        self.assertEqual(summary["deleted"], 1)
        self.assertNotIn("topology/a.md", self.rows_by_path())
        self.assert_equals_full_rebuild()

    def test_resume_is_byte_identical_when_nothing_changed(self):
        self.write("a.md", "# 原理\nBUCK 电感需要根据压差和频率计算。")
        self.write("b.md", "# 计算\n电容纹波与开关频率有关。")
        with self._patch_parse_failure("b.md"):
            with self.assertRaises(RuntimeError):
                self.full()

        summary = self.resume()

        # a.md 在 .tmp 里已完整 -> 复用；b.md 上次就失败了 -> 重做
        self.assertEqual(summary["unchanged"], 1)
        self.assertEqual(summary["added"], 1)
        self.assertEqual(summary["changed"], 0)
        self.assert_equals_full_rebuild()

    def test_resume_refuses_when_slicing_config_changed_after_crash(self):
        # 中断后为了修复问题改了切片配置：续跑会把两种口径混在一起，必须拦下。
        self.write("a.md", "# 原理\nBUCK 电感需要根据压差和频率计算。")
        self.write("b.md", "# 计算\n电容纹波与开关频率有关。")
        with self._patch_parse_failure("b.md"):
            with self.assertRaises(RuntimeError):
                self.full()

        with patch.object(settings.ingestion.chunking, "section_max_chars", self.SIZE + 77):
            with self.assertRaisesRegex(RuntimeError, "切片口径已变化"):
                self.resume()

    def test_incremental_stem_twin_matches_under_unicode_normalization(self):
        # NFD / NFC 两种写法在 NTFS 上是两个不同文件，而指纹侧会做 NFC 归一化
        # （chunk_clean.normalize_text）。分组键必须同口径，否则会出现「指纹判为同一块、
        # 分组判为不同孪生」的错配，被挤掉的块永远无法复活。
        nfd = unicodedata.normalize("NFD", "caf\u00e9") + ".md"
        nfc = unicodedata.normalize("NFC", "caf\u00e9") + ".md"
        shared = "## 共用章节\nBUCK 电感需要根据压差和频率计算。\n"
        self.write(nfd, "# 文档\n" + shared + "## 仅A\nMOSFET 导通损耗按电流估算。\n")
        self.write(nfc, "# 文档\n" + shared + "## 仅B\n环路补偿要保证相位裕度。\n")

        from ingestion.index_state import normalized_stem

        self.assertEqual(normalized_stem("topology/" + nfd), normalized_stem("topology/" + nfc))

        self.full()
        rows = self.rows_by_path()
        self.assertEqual(len(rows["topology/" + nfd]), 2)
        self.assertEqual(len(rows["topology/" + nfc]), 1)  # 共用块被 NFD 那份挤掉

        self.remove(nfd)
        summary = self.incremental()

        self.assertEqual(summary["changed"], 1)  # NFC 那份被孪生扩张提升
        self.assertEqual(len(self.rows_by_path()["topology/" + nfc]), 2)
        self.assert_equals_full_rebuild()

    def test_pipeline_fingerprint_tracks_chunking_tokenizer(self):
        # 切片 tokenizer 取自 embedding 模型（docling_chunker 的 HybridChunker），
        # 换模型必须让指纹变化，否则增量会把两种切分口径混进同一份索引。
        from ingestion import index_state

        before = index_state.pipeline_fingerprint()
        with patch.object(settings.model.embedding, "name", "BAAI/another-model"):
            self.assertNotEqual(index_state.pipeline_fingerprint(), before)

    def test_every_chunk_starts_with_context_header(self):
        # 孪生分组的充分性依赖这个不变量（见 ingestion.index_state.normalized_stem）。
        self.write("a.md", "# 原理\nBUCK 电感需要根据压差和频率计算。")
        (self.kb / "topology" / "b.csv").write_text(
            "Name,Value\nInductor,10 uH\nCapacitor,20 uF\n", encoding="utf-8"
        )
        for name in ("topology/a.md", "topology/b.csv"):
            with self.subTest(target=name):
                chunks = clean_chunk.build_chunks_for_file(self.kb / name, self.kb, self.SIZE)
                self.assertTrue(chunks)
                for chunk in chunks:
                    self.assertTrue(chunk.content.startswith("文档《"), chunk.content[:40])

    # ---------- 保护闸门与 dry-run ----------

    def test_incremental_parser_failure_keeps_published_index(self):
        self.write("a.md", "# 原理\nBUCK 电感需要根据压差和频率计算。")
        self.full()
        before = self.output.read_bytes()

        self.write("b.md", "# 计算\n电容纹波与开关频率有关。")
        with patch.object(clean_chunk, "build_chunks_for_file", side_effect=ValueError("SECRET BODY")):
            with self.assertRaises(RuntimeError):
                self.incremental()

        self.assertEqual(self.output.read_bytes(), before)
        self.assertNotIn("SECRET BODY", self.errors.read_text(encoding="utf-8"))

    def test_incremental_refuses_when_kb_scan_is_empty(self):
        self.write("a.md", "# 原理\nBUCK 电感需要根据压差和频率计算。")
        self.full()
        before = self.output.read_bytes()

        with patch.object(clean_chunk, "scan_supported_files", return_value=[]):
            with self.assertRaisesRegex(RuntimeError, "扫描结果为 0"):
                self.incremental()

        self.assertEqual(self.output.read_bytes(), before)

    def test_incremental_refuses_mass_delete_without_force(self):
        for index in range(4):
            self.write(f"f{index}.md", f"# 第{index}章\nBUCK 电感与频率的关系说明。")
        self.full()

        for index in range(2):
            self.remove(f"f{index}.md")

        with patch.object(settings.ingestion.incremental, "max_delete_ratio", 0.25):
            with self.assertRaisesRegex(RuntimeError, "force-delete"):
                self.incremental()
            summary = self.incremental(force_delete=True)
        self.assertEqual(summary["deleted"], 2)
        self.assert_equals_full_rebuild()

    def test_incremental_dry_run_writes_nothing(self):
        self.write("a.md", "# 原理\nBUCK 电感需要根据压差和频率计算。")
        self.full()
        self.write("b.md", "# 计算\n电容纹波与开关频率有关。")

        watched = (self.kb, self.store, self.output, self.errors, self.state_file())
        before = self.tree_stamp(*watched)
        summary = self.incremental(dry_run=True)

        self.assertEqual(summary["added"], 1)
        self.assertEqual(self.tree_stamp(*watched), before)

    def test_incremental_refuses_when_pipeline_fingerprint_changed(self):
        self.write("a.md", "# 原理\nBUCK 电感需要根据压差和频率计算。")
        self.full()

        with patch.object(settings.ingestion.docling, "max_tokens", 999):
            with self.assertRaisesRegex(RuntimeError, "切片口径已变化"):
                self.incremental()

    def test_incremental_tolerates_fingerprint_coverage_change(self):
        # 指纹的覆盖范围本身会演进（后来新增了 tokenizer、解析器版本等输入项）。
        # 只要两边都有的输入项取值一致，就不该逼用户把几小时的全量重切跑一遍。
        self.write("a.md", "# 原理\nBUCK 电感需要根据压差和频率计算。")
        self.full()

        state = json.loads(self.state_file().read_text(encoding="utf-8"))
        state.pop("fingerprint_parts")        # 模拟旧版 sidecar：只有哈希，没有可比项
        state["pipeline_fingerprint"] = "old-definition"
        self.state_file().write_text(json.dumps(state), encoding="utf-8")

        summary = self.incremental()

        self.assertEqual(summary["unchanged"], 1)
        self.assertEqual(summary["chunks"], 0)
        # 接管来的基线没被校验过，标记为未确认，告警会持续到跑过一次全量。
        written = json.loads(self.state_file().read_text(encoding="utf-8"))
        self.assertFalse(written["fingerprint_verified"])

    def test_incremental_bootstraps_from_existing_index_without_state(self):
        self.write("a.md", "# 原理\nBUCK 电感需要根据压差和频率计算。")
        self.write("b.md", "# 计算\n电容纹波与开关频率有关。")
        self.full()
        before = self.output.read_bytes()

        # sidecar 丢失（切换到增量模式前用旧版本入的库）时，应以现有索引为基线继续。
        self.state_file().unlink()
        summary = self.incremental()

        self.assertEqual(summary["unchanged"], 2)
        self.assertEqual(summary["chunks"], 0)
        self.assertEqual(self.output.read_bytes(), before)

    # ---------- 真实 Chroma ----------

    def _spawn_subprocess(self, pattern, timeout=180):
        result = subprocess.run(
            [sys.executable, "-X", "utf8", "-m", "unittest", "discover",
             "-s", "tests", "-p", "test_index_only.py", "-k", pattern],
            cwd=Path(__file__).resolve().parents[1],
            env={**os.environ, "RAG_TEST_SUBPROCESS_ROOT": str(self.root)},
            capture_output=True, encoding="utf-8", timeout=timeout,
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    @unittest.skipUnless(importlib.util.find_spec("chromadb"), "需要 chromadb")
    def test_real_chroma_incremental_add_change_delete(self):
        if not os.getenv("RAG_TEST_SUBPROCESS_ROOT"):
            self._spawn_subprocess("real_chroma_incremental")
            return

        import chromadb
        import numpy as np
        from ingestion import ingest
        from utils.index_audit import assert_index_matches_index_file

        class Model:
            def encode(self, texts, **kwargs):
                return np.array([[1.0, 0.1, 0.0]] * len(texts))

        self.write("a.md", "# 原理\nBUCK 电感需要根据压差和频率计算。")
        self.write("b.md", "# 计算\n电容纹波与开关频率有关。")
        client = chromadb.PersistentClient(path=str(self.root / "chroma"))
        collection_name = settings.paths.collection_prefix + "topology"

        with patch.object(settings.paths, "chroma_dir", self.root / "chroma"), \
             patch.object(embed_store, "get_client", return_value=client), \
             patch.object(embed_store, "load_embedding_model", return_value=Model()):
            ingest.main([])
            collection = client.get_collection(name=collection_name, embedding_function=None)
            self.assertEqual(collection.count(), len(self.rows_by_path().get("topology/a.md", []))
                             + len(self.rows_by_path().get("topology/b.md", [])))
            assert_index_matches_index_file(client, self.output)

            # 新增一个文件
            self.write("c.md", "# 布局\nPCB 走线宽度影响温升。")
            ingest.main(["--incremental"])
            assert_index_matches_index_file(client, self.output)

            # 修改一个文件
            self.write("a.md", "# 原理\n换过的正文：电感换成电容，频率换成占空比。")
            ingest.main(["--incremental"])
            assert_index_matches_index_file(client, self.output)

            # 删除一个文件：其 chunk 必须从向量库消失
            self.remove("b.md")
            ingest.main(["--incremental"])
            assert_index_matches_index_file(client, self.output)

            removed_ids = {
                json.loads(line)["chunk_id"]
                for line in self.rows_by_path().get("topology/b.md", [])
            }
            self.assertFalse(removed_ids)

            # 对账幂等：再跑一次不应有任何变化
            count_before = collection.count()
            ingest.main([])
            self.assertEqual(collection.count(), count_before)

    @unittest.skipUnless(importlib.util.find_spec("chromadb"), "需要 chromadb")
    def test_real_chroma_rebuild_rewrites_every_vector(self):
        if not os.getenv("RAG_TEST_SUBPROCESS_ROOT"):
            self._spawn_subprocess("real_chroma_rebuild")
            return

        import chromadb
        import numpy as np
        from ingestion import ingest

        class Model:
            calls = 0

            def encode(self, texts, **kwargs):
                Model.calls += 1
                return np.array([[1.0, 0.1, 0.0]] * len(texts))

        self.write("a.md", "# 原理\nBUCK 电感需要根据压差和频率计算。")
        self.write("b.md", "# 计算\n电容纹波与开关频率有关。")
        client = chromadb.PersistentClient(path=str(self.root / "chroma"))

        with patch.object(settings.paths, "chroma_dir", self.root / "chroma"), \
             patch.object(embed_store, "get_client", return_value=client), \
             patch.object(embed_store, "load_embedding_model", return_value=Model()):
            ingest.main([])
            self.assertGreater(Model.calls, 0)
            collection = client.get_collection(
                name=settings.paths.collection_prefix + "topology", embedding_function=None
            )
            total = collection.count()

            # 无变化时的对账应当连模型都不加载
            Model.calls = 0
            summary = embed_store.run(self.output)
            self.assertEqual(Model.calls, 0)
            self.assertEqual(summary["topology"], {"added": 0, "deleted": 0})

            # --rebuild 必须把已在库里的 chunk 全部重写（而不是崩在写集校验上）
            Model.calls = 0
            summary = embed_store.run(self.output, rebuild=True)
            self.assertEqual(Model.calls, 1)
            self.assertEqual(summary["topology"]["added"], total)
            self.assertEqual(collection.count(), total)

            # 换 embedding 模型 -> 自动强制重嵌，且此后不再重复触发
            with patch.object(settings.model.embedding, "name", "BAAI/another-model"):
                Model.calls = 0
                summary = embed_store.run(self.output)
                self.assertEqual(Model.calls, 1)
                self.assertEqual(summary["topology"]["added"], total)

                Model.calls = 0
                embed_store.run(self.output)
                self.assertEqual(Model.calls, 0)

    @unittest.skipUnless(importlib.util.find_spec("chromadb"), "需要 chromadb")
    def test_real_chroma_full_run_prunes_orphans(self):
        if not os.getenv("RAG_TEST_SUBPROCESS_ROOT"):
            self._spawn_subprocess("real_chroma_full_run_prunes")
            return

        import chromadb
        import numpy as np
        from ingestion import ingest
        from utils.index_audit import assert_index_matches_index_file

        class Model:
            def encode(self, texts, **kwargs):
                return np.array([[1.0, 0.1, 0.0]] * len(texts))

        self.write("a.md", "# 原理\nBUCK 电感需要根据压差和频率计算。")
        self.write("b.md", "# 计算\n电容纹波与开关频率有关。")
        client = chromadb.PersistentClient(path=str(self.root / "chroma"))
        collection_name = settings.paths.collection_prefix + "topology"

        with patch.object(settings.paths, "chroma_dir", self.root / "chroma"), \
             patch.object(embed_store, "get_client", return_value=client), \
             patch.object(embed_store, "load_embedding_model", return_value=Model()):
            ingest.main([])
            orphan_ids = [json.loads(line)["chunk_id"] for line in self.rows_by_path()["topology/b.md"]]
            self.assertTrue(orphan_ids)

            # 全量重跑同样必须清掉被删文件的向量（对账器与增量共用一条写入路径）。
            self.remove("b.md")
            ingest.main([])

            collection = client.get_collection(name=collection_name, embedding_function=None)
            remaining = set(collection.get(include=[])["ids"])
            self.assertFalse(remaining & set(orphan_ids))
            assert_index_matches_index_file(client, self.output)


if __name__ == "__main__":
    unittest.main()
