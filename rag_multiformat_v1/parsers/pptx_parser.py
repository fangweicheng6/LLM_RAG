# -*- coding: utf-8 -*-
from __future__ import annotations

import tempfile
from pathlib import Path
from typing import List

from schema import CanonicalBlock
from parsers.base import BaseParser
from parsers.ocr_utils import ocr_text


def _table_to_markdown(table) -> str:
    rows = []

    for row in table.rows:
        rows.append(
            [
                str(cell.text or "").strip().replace("\n", " ").replace("|", r"\|")
                for cell in row.cells
            ]
        )

    if not rows:
        return ""

    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]

    lines = [
        "| " + " | ".join(rows[0]) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]

    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")

    return "\n".join(lines)


class PPTXParser(BaseParser):
    extensions = (".pptx",)
    document_type = "powerpoint"

    def parse(self, path: Path) -> List[CanonicalBlock]:
        try:
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE
        except ImportError as e:
            raise RuntimeError(
                "解析 PPTX 需要 python-pptx。\n"
                "请执行：pip install python-pptx"
            ) from e

        prs = Presentation(path)
        blocks: List[CanonicalBlock] = []

        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_title = ""

            title_shape = slide.shapes.title
            if title_shape is not None and getattr(title_shape, "text", "").strip():
                slide_title = title_shape.text.strip()

            section = f"Slide {slide_idx}"
            if slide_title:
                section += f" > {slide_title}"

            image_index = 0

            # 按 top / left 排序，尽量接近视觉阅读顺序。
            shapes = sorted(
                list(slide.shapes),
                key=lambda s: (
                    int(getattr(s, "top", 0)),
                    int(getattr(s, "left", 0)),
                ),
            )

            for shape in shapes:
                if shape is title_shape:
                    continue

                if getattr(shape, "has_table", False):
                    md = _table_to_markdown(shape.table)

                    if md:
                        blocks.append(
                            CanonicalBlock(
                                content_type="table",
                                content=md,
                                section_path=section,
                                slide_number=slide_idx,
                            )
                        )
                    continue

                if getattr(shape, "has_text_frame", False):
                    text = (getattr(shape, "text", "") or "").strip()

                    if text:
                        blocks.append(
                            CanonicalBlock(
                                content_type="slide_text",
                                content=text,
                                section_path=section,
                                slide_number=slide_idx,
                            )
                        )
                    continue

                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    image_index += 1

                    try:
                        image = shape.image
                        suffix = "." + (image.ext or "png")

                        with tempfile.NamedTemporaryFile(
                            suffix=suffix,
                            delete=False,
                        ) as tmp:
                            tmp_path = Path(tmp.name)
                            tmp.write(image.blob)

                        try:
                            text = ocr_text(tmp_path)
                        finally:
                            tmp_path.unlink(missing_ok=True)

                        if text:
                            blocks.append(
                                CanonicalBlock(
                                    content_type="image_ocr",
                                    content=text,
                                    section_path=section,
                                    slide_number=slide_idx,
                                    image_index=image_index,
                                )
                            )

                    except Exception:
                        continue

        return blocks
