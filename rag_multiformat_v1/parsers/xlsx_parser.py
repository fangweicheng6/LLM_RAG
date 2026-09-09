# -*- coding: utf-8 -*-
from __future__ import annotations

from pathlib import Path
from typing import List

from schema import CanonicalBlock
from parsers.base import BaseParser
from parsers.common import row_to_retrieval_text


class XLSXParser(BaseParser):
    extensions = (".xlsx", ".xlsm")
    document_type = "spreadsheet"

    def parse(self, path: Path) -> List[CanonicalBlock]:
        try:
            from openpyxl import load_workbook
        except ImportError as e:
            raise RuntimeError(
                "解析 Excel 需要 openpyxl。\n"
                "请执行：pip install openpyxl"
            ) from e

        wb = load_workbook(path, read_only=True, data_only=True)
        blocks: List[CanonicalBlock] = []

        try:
            for ws in wb.worksheets:
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    continue

                header_idx = None
                headers = None

                for idx, row in enumerate(rows, start=1):
                    if any(v is not None and str(v).strip() != "" for v in row):
                        header_idx = idx
                        headers = list(row)
                        break

                if headers is None:
                    continue

                for row_idx, row in enumerate(rows[header_idx:], start=header_idx + 1):
                    if not any(v is not None and str(v).strip() != "" for v in row):
                        continue

                    content = row_to_retrieval_text(headers, list(row))
                    if not content:
                        continue

                    blocks.append(
                        CanonicalBlock(
                            content_type="table_row",
                            content=content,
                            section_path=ws.title,
                            sheet_name=ws.title,
                            row_start=row_idx,
                            row_end=row_idx,
                        )
                    )

        finally:
            wb.close()

        return blocks
