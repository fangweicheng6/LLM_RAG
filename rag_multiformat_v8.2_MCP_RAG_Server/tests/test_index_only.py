"""无需模型或 API key；真实 Chroma 测试需要安装 chromadb。

python -m unittest discover -s tests -v
"""
import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import ingestion.clean_chunk as clean_chunk
import ingestion.embed_store as embed_store
import ingestion.source_store as source_store
from utils.config_loader import settings
from utils.index_audit import assert_index_collection, validate_storage_paths
from utils.schema import CanonicalBlock
from ingestion.source_store import SourceResolver, SourceResolutionError, located_blocks


class SourceTests(unittest.TestCase):
    def setUp(self):
        child_root = os.getenv("RAG_TEST_SUBPROCESS_ROOT")
        if child_root:
            self.root = Path(child_root)
        else:
            self.temp = tempfile.TemporaryDirectory()
            self.addCleanup(self.temp.cleanup)
            self.root = Path(self.temp.name)
        self.kb = self.root / "kb"
        (self.kb / "topology").mkdir(parents=True, exist_ok=True)
        self.store = self.root / "source"
        self.patch = patch.object(settings.paths, "source_store_dir", self.store)
        self.patch.start()
        self.addCleanup(self.patch.stop)

    def build(self, name="guide.txt", text="BUCK 的电感纹波需要根据开关频率计算。\n完整原文不能出现在向量数据库里。", size=40):
        path = self.kb / "topology" / name
        path.write_text(text, encoding="utf-8")
        return clean_chunk.build_chunks_for_file(path, self.kb, size)

    def test_jsonl_has_only_index_and_roundtrips_full_content(self):
        chunks = self.build()
        output = self.root / "index.jsonl"
        output.write_text("\n".join(json.dumps(c.to_index_dict(), ensure_ascii=False) for c in chunks), encoding="utf-8")
        loaded = embed_store.load_chunks(output)
        self.assertEqual([c.content for c in loaded], [c.content for c in chunks])
        for chunk in chunks:
            index = chunk.to_index_dict()
            self.assertFalse({"content", "document", "documents", "title", "section_path", "tags"} & set(index))
            self.assertTrue(json.loads(index["keywords"]))
            item = SourceResolver().hydrate({"chunk_id": chunk.chunk_id, "metadata": chunk.chroma_metadata()})
            self.assertEqual(item["document"], chunk.content)
            self.assertEqual(item["metadata"]["section_path"], chunk.section_path)

    def test_lines_columns_pages_and_repeated_text_are_exact(self):
        blocks = [CanonicalBlock("text", "第一行\n" + "A" * 27 + "\n重复行\n重复行", section_path="同一章节", page=1),
                  CanonicalBlock("text", "第二页\n最后一行", section_path="同一章节", page=2)]
        units, slices = located_blocks(blocks, "title", 11)
        self.assertEqual(len(units), 2)
        for unit_id, text in units.items():
            self.assertEqual("".join(block.content for block, loc in slices if loc["source_unit"] == unit_id), text)
        for block, loc in slices:
            text = units[loc["source_unit"]]
            self.assertEqual(block.content, text[loc["char_start"]:loc["char_end"]])
            self.assertLessEqual(len(block.content), 11)
            self.assertEqual(loc["line_start"], text[:loc["char_start"]].count("\n") + 1)
            self.assertEqual(loc["line_end"], text[:loc["char_end"] - 1].count("\n") + 1)
        self.assertTrue(any(loc["column_start"] > 1 for _, loc in slices))
        self.assertEqual({block.page for block, _ in slices}, {1, 2})

    def test_all_format_location_fields_survive(self):
        blocks = [CanonicalBlock("table_row", "单元格1\n单元格2", sheet_name="Sheet A", row_start=3, row_end=3),
                  CanonicalBlock("slide_text", "幻灯片完整内容", slide_number=2),
                  CanonicalBlock("image_ocr", "图片文字完整内容", image_index=1),
                  CanonicalBlock("transcript", "音频转录完整内容", media_start_sec=2.0, media_end_sec=5.0)]
        _, slices = located_blocks(blocks, "title", 100)
        for original, (result, _) in zip(blocks, slices):
            for field in ("sheet_name", "row_start", "row_end", "slide_number", "image_index", "media_start_sec", "media_end_sec"):
                self.assertEqual(getattr(original, field), getattr(result, field))

    def test_tampered_missing_or_wrong_locator_fails_closed(self):
        chunk = self.build()[0]
        item = {"chunk_id": chunk.chunk_id, "metadata": chunk.chroma_metadata()}
        for key, value in (("char_start", 999), ("domain", "device"), ("line_start", 999), ("source_ref", "../../secret")):
            with self.subTest(key=key), self.assertRaises(SourceResolutionError):
                SourceResolver().hydrate({**item, "metadata": {**item["metadata"], key: value}})
        snapshot = self.store / chunk.source_ref[:2] / (chunk.source_ref + ".json")
        snapshot.write_text("tampered", encoding="utf-8")
        with self.assertRaises(SourceResolutionError):
            SourceResolver().hydrate(item)
        snapshot.unlink()
        with self.assertRaises(SourceResolutionError):
            SourceResolver().hydrate(item)

    def test_source_versions_do_not_silently_change_old_results(self):
        old = self.build()[0]
        new = self.build(text="更新之后的原文已经完全不同。\n只有重建索引才应检索新版本。")[0]
        self.assertNotEqual(old.source_ref, new.source_ref)
        for chunk in (old, new):
            self.assertEqual(SourceResolver().hydrate({"chunk_id": chunk.chunk_id, "metadata": chunk.chroma_metadata()})["document"], chunk.content)

    def test_short_content_is_not_lost(self):
        # Docling 路径会前缀确定性上下文头（CCH）；短正文必须完整保留在正文部分。
        content = self.build(text="短句")[0].content
        self.assertTrue(content.endswith("短句"))

    def test_old_plaintext_jsonl_is_rejected(self):
        path = self.root / "old.jsonl"
        path.write_text(json.dumps(self.build()[0].to_dict()), encoding="utf-8")
        with self.assertRaisesRegex(RuntimeError, "旧版正文"):
            embed_store.load_chunks(path)

    def test_parser_failure_preserves_previous_manifest_and_does_not_log_body(self):
        self.build()
        output, error = self.root / "index.jsonl", self.root / "errors.jsonl"
        output.write_text("previous index", encoding="utf-8")
        with patch.object(clean_chunk, "build_chunks_for_file", side_effect=ValueError("SECRET BODY")):
            with self.assertRaises(RuntimeError):
                clean_chunk.run(self.kb, output, error, 100)
        self.assertEqual(output.read_text(encoding="utf-8"), "previous index")
        self.assertNotIn("SECRET BODY", error.read_text(encoding="utf-8"))

    def test_source_store_must_be_outside_chroma(self):
        for chroma, source in ((self.root, self.root / "source"), (self.root / "db", self.root), (self.root, self.root)):
            with self.assertRaises(ValueError):
                validate_storage_paths(chroma, source)
        validate_storage_paths(self.root / "db", self.store)

    @unittest.skipUnless(importlib.util.find_spec("chromadb"), "需要 chromadb")
    def test_real_chroma_retrieval_expansion_reingest_and_audit(self):
        # Chroma 的 Rust/SQLite 句柄在 Windows 上可能持有到进程退出。
        # 子进程跑真实持久库断言，退出后父进程再删除临时数据。
        if not os.getenv("RAG_TEST_SUBPROCESS_ROOT"):
            result = subprocess.run(
                [sys.executable, "-X", "utf8", "-m", "unittest", "discover", "-s", "tests", "-k", "real_chroma"],
                cwd=Path(__file__).resolve().parents[1],
                env={**os.environ, "RAG_TEST_SUBPROCESS_ROOT": str(self.root)},
                capture_output=True, encoding="utf-8", timeout=90,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            return
        import chromadb
        import numpy as np
        import search
        import context_expander as expander
        from pipeline.generation.generator import _build_grounded_context

        text = "# 原理\nBUCK parent unique sentence.\n## 计算\nBUCK 电感需要根据压差和频率计算。\n" + "完整长行原文不应存入向量库。" * 9
        chunks = self.build("guide.md", text, size=45)
        client = chromadb.PersistentClient(path=str(self.root / "chroma"))
        collection = embed_store.get_collection(client, "topology")
        vectors = np.array([[1.0, 0.1, 0.0]] * len(chunks))
        embed_store.replace_domain_documents(collection, chunks, vectors)
        self.assertEqual(assert_index_collection(collection, scan=True), len(chunks))
        raw = collection.get(include=["documents", "metadatas", "embeddings"])
        self.assertTrue(all(d is None for d in raw["documents"]))
        self.assertEqual(len(raw["embeddings"]), len(chunks))
        self.assertNotIn("完整长行原文不应存入向量库。", json.dumps(raw["metadatas"], ensure_ascii=False))

        found = search.query_collection(collection=collection, query_embedding=vectors[0], candidate_limit=20)
        self.assertEqual({i["document"] for i in found}, {c.content for c in chunks})
        anchor = next(i for i in found if ">" in i["metadata"]["section_path"])
        section = expander.expand_same_section(client, anchor)
        self.assertTrue(section)
        self.assertTrue(all(i["document"] for i in section))
        parent = expander.expand_parent_section(client, anchor)
        self.assertTrue(parent)
        self.assertIn("parent unique", parent[0]["document"])
        neighbors = expander.expand_neighbors(client, anchor, before=1, after=1)
        self.assertTrue(neighbors)
        class Model:
            def encode(self, texts, **kwargs):
                return np.array([[1.0, 0.1, 0.0]] * len(texts))
        targeted = expander.expand_targeted_search(client=client, embedding_model=Model(), query="BUCK", domains=["topology"])
        self.assertTrue(targeted)
        self.assertTrue(all(i["document"] for i in targeted))
        prompt, sources = _build_grounded_context(found)
        self.assertIn(found[0]["document"], prompt)
        self.assertEqual(len(sources), len(found))

        # 完整入口联调：真实解析/快照/索引/Chroma；替换模型避免下载和外部调用。
        import ingest
        manifest = self.root / "processed" / "index.jsonl"
        with patch.object(settings.paths, "kb_root", self.kb), \
             patch.object(settings.paths, "chunk_jsonl", manifest), \
             patch.object(settings.paths, "error_jsonl", self.root / "errors.jsonl"), \
             patch.object(settings.ingestion.chunking, "section_max_chars", 45), \
             patch.object(settings.paths, "chroma_dir", self.root / "chroma"), \
             patch.object(settings.paths, "source_store_dir", self.store), \
             patch.object(embed_store, "get_client", return_value=client), \
             patch.object(embed_store, "load_embedding_model", return_value=Model()):
            ingest.main()
        self.assertNotIn('"content":', manifest.read_text(encoding="utf-8"))
        self.assertEqual(collection.count(), len(chunks))
        def choose(**kwargs):
            return {"sufficient": True, "reason": "test", "expansion_requests": []}, kwargs["candidates"]
        def answer(**kwargs):
            content, references = _build_grounded_context(kwargs["selected_contexts"])
            return content, references
        with patch.object(search, "get_client", return_value=client), \
             patch.object(search, "load_embedding_model", return_value=Model()), \
             patch.object(search, "create_llm_client", return_value=object()), \
             patch.object(search, "analyze_query", return_value={"domains": ["topology"], "keywords": ["BUCK"]}), \
             patch.object(search, "select_contexts", side_effect=choose), \
             patch.object(search, "generate_answer", side_effect=answer):
            response = search.run("BUCK calculation")
        self.assertTrue(response["selected_contexts"])
        self.assertIn("lines=", response["answer_sources"][0]["location"])

        # 替换同文档后旧分块必须消失，当前记录依然无正文。
        updated = self.build("guide.md", "# 原理\n更新的正文只在本地原文目录中。", size=45)
        embed_store.replace_domain_documents(collection, updated, vectors[:len(updated)])
        self.assertEqual(collection.count(), len(updated))
        self.assertEqual(assert_index_collection(collection, scan=True), len(updated))
        embed_store.replace_domain_documents(collection, updated, vectors[:len(updated)])
        self.assertEqual(collection.count(), len(updated))

        # 有版本标记也不能绕过正文审计。
        collection.upsert(ids=["bad"], documents=["SECRET BODY"], embeddings=[[1., 0., 0.]], metadatas=[updated[0].chroma_metadata()])
        with self.assertRaisesRegex(RuntimeError, "documents"):
            assert_index_collection(collection, scan=True)
        with self.assertRaises(RuntimeError):
            embed_store.replace_domain_documents(collection, updated, vectors[:len(updated)])
        legacy = client.create_collection("legacy", embedding_function=None)
        with self.assertRaisesRegex(RuntimeError, "旧版"):
            assert_index_collection(legacy)

    @unittest.skipUnless(all(importlib.util.find_spec(m) for m in ("docx", "pptx", "openpyxl", "bs4", "lxml")), "需要 Office/HTML 解析器依赖")
    def test_real_office_html_csv_files_roundtrip(self):
        from docx import Document
        from pptx import Presentation
        from openpyxl import Workbook

        directory = self.kb / "topology"
        doc = Document()
        doc.add_heading("Design", level=1)
        doc.add_paragraph("BUCK original paragraph with formula L = 10 uH.")
        doc.save(directory / "guide.docx")
        slides = Presentation()
        slide = slides.slides.add_slide(slides.slide_layouts[1])
        slide.shapes.title.text = "BUCK Design"
        slide.placeholders[1].text = "Original slide content and frequency details."
        slides.save(directory / "guide.pptx")
        book = Workbook()
        book.active.title = "Components"
        book.active.append(["Name", "Value"])
        book.active.append(["Inductor", "10 uH"])
        book.save(directory / "guide.xlsx")
        book.close()
        (directory / "guide.csv").write_text("Name,Value\nInductor,10 uH\nCapacitor,20 uF\n", encoding="utf-8")
        (directory / "guide.html").write_text("<h1>BUCK</h1><p>Original HTML paragraph with frequency details.</p>", encoding="utf-8")
        for path in directory.iterdir():
            with self.subTest(format=path.suffix):
                chunks = clean_chunk.build_chunks_for_file(path, self.kb, 30)
                self.assertTrue(chunks)
                for chunk in chunks:
                    restored = SourceResolver().hydrate({"chunk_id": chunk.chunk_id, "metadata": chunk.chroma_metadata()})
                    self.assertEqual(restored["document"], chunk.content)
                    self.assertIsNone(chunk.page)
                    if path.suffix in (".xlsx", ".csv"):
                        self.assertGreaterEqual(chunk.row_start, 2)
                        self.assertEqual(chunk.row_start, chunk.row_end)
                    if path.suffix == ".pptx":
                        self.assertEqual(chunk.slide_number, 1)

    @unittest.skipUnless(importlib.util.find_spec("fitz"), "需要 PyMuPDF")
    def test_real_pdf_pages_remain_separate(self):
        import fitz
        path = self.kb / "topology" / "two_pages.pdf"
        doc = fitz.open()
        for number in (1, 2):
            page = doc.new_page()
            page.insert_text((72, 72), f"Page {number} first source line.\nSecond source line with BUCK calculation and frequency details.")
        doc.save(path)
        doc.close()
        chunks = clean_chunk.build_chunks_for_file(path, self.kb, 1000)
        self.assertEqual({c.page for c in chunks}, {1, 2})
        for chunk in chunks:
            self.assertEqual(chunk.line_start, 1)
            self.assertEqual(chunk.line_end, 2)
            self.assertEqual(SourceResolver().hydrate({"chunk_id": chunk.chunk_id, "metadata": chunk.chroma_metadata()})["document"], chunk.content)


if __name__ == "__main__":
    unittest.main()
